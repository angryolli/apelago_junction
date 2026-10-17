"""Build the Ocean Opportunity Radar pitch-deck template."""

from __future__ import annotations

import io
from pathlib import Path

from lxml import etree
from PIL import Image, ImageEnhance
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
LOGO = ROOT / "frontend/fff13e3f-08a7-4f29-a5c3-4a45fe9371c1/public/logo.png"
CLOSING = Path(__file__).resolve().parent / "assets" / "closing.jpg"
OUT = Path(__file__).resolve().parent / "Ocean-Opportunity-Radar.pptx"

W = 13.333333333333334
H = 7.5
ML = 0.56
CONTENT_W = W - ML * 2

NAVY = RGBColor(
    0x0B,
    0x3A,
    0x54,
)
INK = RGBColor(
    0x14,
    0x32,
    0x42,
)
TEAL = RGBColor(
    0x0E,
    0x8A,
    0x90,
)
TEAL_DEEP = RGBColor(
    0x0A,
    0x6A,
    0x78,
)
TEAL_BRIGHT = RGBColor(
    0x1A,
    0xB5,
    0xB0,
)
ICE = RGBColor(
    0xF4,
    0xFB,
    0xFA,
)
WHITE = RGBColor(
    0xFF,
    0xFF,
    0xFF,
)
HINT = RGBColor(
    0x5C,
    0x78,
    0x86,
)
LINE = RGBColor(
    0xD3,
    0xEB,
    0xE9,
)
SOFT = RGBColor(
    0xE7,
    0xF6,
    0xF4,
)
MIST = RGBColor(
    0xD7,
    0xF0,
    0xEE,
)

DISPLAY = "Avenir Next"
BODY = "Avenir Next"


def knock_out_white(image: Image.Image, threshold: int = 247) -> Image.Image:
    """Make the logo's white field transparent, keeping the pale radar rings."""
    image = image.convert("RGBA")
    pixels = image.load()
    width, height = image.size

    for y in range(height):
        for x in range(width):
            red, green, blue, alpha = pixels[x, y]
            if red >= threshold and green >= threshold and blue >= threshold:
                pixels[x, y] = (red, green, blue, 0)

    return image


def cropped_logo() -> io.BytesIO:
    """Return the full lockup with the white field removed."""
    image = knock_out_white(Image.open(LOGO))
    bbox = image.getchannel("A").getbbox()
    image = image.crop(bbox)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    buffer.aspect = image.size[0] / image.size[1]

    return buffer


def icon_mark() -> io.BytesIO:
    """Return the lighthouse mark alone, for slide headers."""
    image = knock_out_white(Image.open(LOGO))
    image = image.crop((0, 0, 368, image.height))
    bbox = image.getchannel("A").getbbox()
    image = image.crop(bbox)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    buffer.aspect = image.size[0] / image.size[1]

    return buffer


def force_typeface(run, name: str) -> None:
    """Pin latin, east-asian, and complex-script faces to one family."""
    props = run._r.get_or_add_rPr()

    for tag in ("a:latin", "a:ea", "a:cs"):
        node = props.find(qn(tag))

        if node is None:
            node = etree.SubElement(props, qn(tag))

        node.set("typeface", name)


def set_spacing(run, hundredths: int) -> None:
    """Set character spacing in hundredths of a point."""
    run._r.get_or_add_rPr().set("spc", str(hundredths))


def fill_rgb(
    shape,
    color: RGBColor,
    alpha: int | None = None,
) -> None:
    """Apply a solid fill, with optional alpha in OOXML units (100000 = opaque)."""
    shape.fill.solid()
    shape.fill.fore_color.rgb = color

    if alpha is None:
        return

    solid = shape._element.spPr.find(qn("a:solidFill"))
    srgb = solid.find(qn("a:srgbClr"))
    etree.SubElement(
        srgb,
        qn("a:alpha"),
        {"val": str(alpha)},
    )


def no_line(shape) -> None:
    shape.line.fill.background()


def add_shadow(
    shape,
    alpha: int = 14000,
    blur: int = 430000,
    dist: int = 70000,
) -> None:
    """Add a soft navy shadow under a card."""
    effect = etree.SubElement(shape._element.spPr, qn("a:effectLst"))
    shadow = etree.SubElement(
        effect,
        qn("a:outerShdw"),
        {
            "blurRad": str(blur),
            "dist": str(dist),
            "dir": "2700000",
            "algn": "tl",
            "rotWithShape": "0",
        },
    )
    color = etree.SubElement(
        shadow,
        qn("a:srgbClr"),
        {"val": "0B3A54"},
    )
    etree.SubElement(
        color,
        qn("a:alpha"),
        {"val": str(alpha)},
    )


def gradient(
    shape,
    left: RGBColor,
    right: RGBColor,
    angle: float = 0,
) -> None:
    shape.fill.gradient()
    shape.fill.gradient_angle = angle
    stops = shape.fill.gradient_stops
    stops[0].position = 0
    stops[0].color.rgb = left
    stops[1].position = 1
    stops[1].color.rgb = right
    no_line(shape)


def rounded(
    slide,
    x,
    y,
    w,
    h,
    radius: float = 0.12,
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    shape.adjustments[0] = radius

    return shape


def write(
    shape,
    paragraphs: list[dict],
    *,
    anchor=MSO_ANCHOR.TOP,
    ml: float = 0.18,
    mr: float = 0.18,
    mt: float = 0.1,
    mb: float = 0.08,
) -> None:
    """Fill a shape text frame from paragraph specs."""
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = None
    frame.margin_left = Inches(ml)
    frame.margin_right = Inches(mr)
    frame.margin_top = Inches(mt)
    frame.margin_bottom = Inches(mb)
    frame._txBody.bodyPr.set(
        "anchor",
        {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor],
    )

    for index, spec in enumerate(paragraphs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = spec.get("align", PP_ALIGN.LEFT)
        paragraph.space_before = Pt(spec.get("before", 0))
        paragraph.space_after = Pt(spec.get("after", 0))

        if "line" in spec:
            paragraph.line_spacing = spec["line"]

        run = paragraph.add_run()
        run.text = spec["text"]
        font = run.font
        font.name = spec.get("font", BODY)
        font.size = Pt(spec["size"])
        font.bold = spec.get("bold", False)
        font.italic = spec.get("italic", False)
        font.color.rgb = spec.get("color", INK)
        force_typeface(run, spec.get("font", BODY))

        if "spc" in spec:
            set_spacing(run, spec["spc"])


def textbox(
    slide,
    x,
    y,
    w,
    h,
    paragraphs,
    **kwargs,
):
    shape = slide.shapes.add_textbox(
        Inches(x),
        Inches(y),
        Inches(w),
        Inches(h),
    )
    write(
        shape,
        paragraphs,
        **kwargs,
    )

    return shape


def blank_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = ICE

    return slide


def notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def header(
    slide,
    mark: io.BytesIO,
    number: str,
    title: str,
    subtitle: str,
) -> None:
    """Shared content-slide header: mark, index, title, and a fill-in line."""
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        0,
        0,
        Inches(W),
        Inches(0.08),
    )
    gradient(
        bar,
        TEAL_BRIGHT,
        TEAL_DEEP,
        angle=0,
    )
    mark.seek(0)
    height = 0.52
    slide.shapes.add_picture(
        mark,
        Inches(ML),
        Inches(0.26),
        Inches(height * mark.aspect),
        Inches(height),
    )
    textbox(
        slide,
        ML + 0.68,
        0.2,
        8,
        0.24,
        [
            {
                "text": number,
                "size": 11,
                "bold": True,
                "color": TEAL,
                "font": DISPLAY,
                "spc": 180,
            },
        ],
        ml=0,
        mr=0,
        mt=0,
        mb=0,
    )
    textbox(
        slide,
        ML + 0.68,
        0.4,
        10,
        0.46,
        [
            {
                "text": title,
                "size": 28,
                "bold": True,
                "color": NAVY,
                "font": DISPLAY,
            },
        ],
        ml=0,
        mr=0,
        mt=0,
        mb=0,
    )
    line = textbox(
        slide,
        ML + 0.68,
        0.88,
        CONTENT_W - 0.68,
        0.3,
        [
            {
                "text": subtitle,
                "size": 14,
                "italic": True,
                "color": HINT,
            },
        ],
        ml=0,
        mr=0,
        mt=0,
        mb=0,
    )
    line.name = f"{title} framing line"
    rule = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(ML),
        Inches(1.26),
        Inches(CONTENT_W),
        Inches(0.012),
    )
    fill_rgb(rule, LINE)
    no_line(rule)


def footer(slide, page: str) -> None:
    rule = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(ML),
        Inches(7.12),
        Inches(CONTENT_W),
        Inches(0.01),
    )
    fill_rgb(rule, LINE)
    no_line(rule)
    textbox(
        slide,
        ML,
        7.16,
        8.2,
        0.26,
        [
            {
                "text": "Ocean Opportunity Radar   ·   A’Pelago  ×  Redstone",
                "size": 11,
                "color": HINT,
                "font": DISPLAY,
            },
        ],
        ml=0,
        mr=0,
        mt=0,
        mb=0,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    textbox(
        slide,
        W - ML - 1.3,
        7.16,
        1.3,
        0.26,
        [
            {
                "text": page,
                "size": 11,
                "bold": True,
                "color": TEAL,
                "font": DISPLAY,
                "align": PP_ALIGN.RIGHT,
                "spc": 80,
            },
        ],
        ml=0,
        mr=0,
        mt=0,
        mb=0,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def card(
    slide,
    x,
    y,
    w,
    h,
    radius: float = 0.1,
):
    shape = rounded(
        slide,
        x,
        y,
        w,
        h,
        radius,
    )
    fill_rgb(shape, WHITE)
    shape.line.color.rgb = LINE
    shape.line.width = Pt(1)
    add_shadow(shape)

    return shape


def pill(
    slide,
    x,
    y,
    label: str,
    size: float = 0.36,
) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(x),
        Inches(y),
        Inches(size),
        Inches(size),
    )
    gradient(
        shape,
        TEAL_BRIGHT,
        TEAL_DEEP,
        angle=135,
    )
    write(
        shape,
        [
            {
                "text": label,
                "size": 12,
                "bold": True,
                "color": WHITE,
                "font": DISPLAY,
                "align": PP_ALIGN.CENTER,
            },
        ],
        anchor=MSO_ANCHOR.MIDDLE,
        ml=0,
        mr=0,
        mt=0,
        mb=0,
    )


def cover(prs: Presentation, lockup: io.BytesIO) -> None:
    slide = blank_slide(prs)
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        0,
        0,
        Inches(W),
        Inches(0.1),
    )
    gradient(
        bar,
        TEAL_BRIGHT,
        TEAL_DEEP,
        angle=0,
    )
    wave_back = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(-1.4),
        Inches(6.55),
        Inches(9.2),
        Inches(2.4),
    )
    fill_rgb(
        wave_back,
        TEAL_BRIGHT,
        alpha=28000,
    )
    no_line(wave_back)
    wave_front = slide.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(5.4),
        Inches(6.78),
        Inches(9.4),
        Inches(2.0),
    )
    fill_rgb(
        wave_front,
        TEAL,
        alpha=36000,
    )
    no_line(wave_front)
    lockup.seek(0)
    width = 7.15
    height = width / lockup.aspect
    slide.shapes.add_picture(
        lockup,
        Inches((W - width) / 2),
        Inches(0.72),
        Inches(width),
        Inches(height),
    )
    rule = rounded(
        slide,
        (W - 1.15) / 2,
        4.28,
        1.15,
        0.045,
        radius=0.5,
    )
    fill_rgb(rule, TEAL)
    no_line(rule)
    textbox(
        slide,
        1.6,
        4.46,
        W - 3.2,
        0.28,
        [
            {
                "text": "PITCH DECK",
                "size": 12,
                "bold": True,
                "color": TEAL,
                "font": DISPLAY,
                "align": PP_ALIGN.CENTER,
                "spc": 280,
            },
        ],
        ml=0,
        mr=0,
        mt=0,
        mb=0,
    )
    pitch = textbox(
        slide,
        1.8,
        4.78,
        W - 3.6,
        0.46,
        [
            {
                "text": "Add the one-line pitch",
                "size": 20,
                "italic": True,
                "color": HINT,
                "align": PP_ALIGN.CENTER,
            },
        ],
        ml=0,
        mr=0,
        mt=0,
        mb=0,
    )
    pitch.name = "One-line pitch"
    labels = ("Presented by", "Audience", "Date")
    prompts = ("Name", "Who this is for", "Day, month, year")
    field_w = 3.15
    gap = 0.36
    total = field_w * 3 + gap * 2
    start = (W - total) / 2

    for index, (label, prompt) in enumerate(zip(labels, prompts)):
        x = start + index * (field_w + gap)
        textbox(
            slide,
            x,
            5.5,
            field_w,
            0.24,
            [
                {
                    "text": label.upper(),
                    "size": 11,
                    "bold": True,
                    "color": TEAL,
                    "font": DISPLAY,
                    "align": PP_ALIGN.CENTER,
                    "spc": 140,
                },
            ],
            ml=0,
            mr=0,
            mt=0,
            mb=0,
        )
        field = textbox(
            slide,
            x,
            5.76,
            field_w,
            0.38,
            [
                {
                    "text": prompt,
                    "size": 16,
                    "italic": True,
                    "color": HINT,
                    "align": PP_ALIGN.CENTER,
                },
            ],
            ml=0,
            mr=0,
            mt=0,
            mb=0,
        )
        field.name = label
        underline = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(x + 0.35),
            Inches(6.18),
            Inches(field_w - 0.7),
            Inches(0.015),
        )
        fill_rgb(
            underline,
            TEAL,
            alpha=70000,
        )
        no_line(underline)

    notes(
        slide,
        "Cover. Replace the one-line pitch, presenter, audience, and date. The logo stays.",
    )


def problem(prs: Presentation, mark: io.BytesIO) -> None:
    slide = blank_slide(prs)
    header(
        slide,
        mark,
        "02",
        "The problem",
        "Add a line that frames the problem.",
    )
    intro = card(
        slide,
        ML,
        1.48,
        CONTENT_W,
        1.42,
        radius=0.08,
    )
    intro.name = "Problem summary"
    write(
        intro,
        [
            {
                "text": "THE GAP",
                "size": 11,
                "bold": True,
                "color": TEAL,
                "font": DISPLAY,
                "spc": 160,
            },
            {
                "text": "Write the problem in two sentences. Who is stuck, and what do they do today?",
                "size": 18,
                "italic": True,
                "color": HINT,
                "before": 8,
                "line": 1.15,
            },
        ],
        mt=0.2,
        ml=0.28,
        mr=0.28,
    )
    titles = ("Scattered sources", "Slow judgement", "Harder in the ocean")
    prompts = (
        "What people have to piece together by hand.",
        "What takes too long to compare or trust.",
        "Why blue-economy startups are easier to miss.",
    )
    gap = 0.2
    card_w = (CONTENT_W - gap * 2) / 3

    for index, (title, prompt) in enumerate(zip(titles, prompts)):
        x = ML + index * (card_w + gap)
        shape = card(
            slide,
            x,
            3.1,
            card_w,
            3.78,
        )
        shape.name = title
        pill(
            slide,
            x + 0.22,
            3.32,
            f"{index + 1:02d}",
        )
        textbox(
            slide,
            x + 0.68,
            3.34,
            card_w - 0.88,
            0.4,
            [
                {
                    "text": title,
                    "size": 16,
                    "bold": True,
                    "color": NAVY,
                    "font": DISPLAY,
                },
            ],
            ml=0,
            mr=0.1,
            mt=0,
            mb=0,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        textbox(
            slide,
            x + 0.16,
            4.0,
            card_w - 0.32,
            2.55,
            [
                {
                    "text": prompt,
                    "size": 15,
                    "italic": True,
                    "color": HINT,
                    "line": 1.2,
                },
            ],
            ml=0.08,
            mr=0.08,
            mt=0,
            mb=0,
        )

    footer(slide, "02  /  06")
    notes(
        slide,
        "Problem. Replace the summary and the three cards. Keep the titles or rewrite them.",
    )


def questions(prs: Presentation, mark: io.BytesIO) -> None:
    slide = blank_slide(prs)
    header(
        slide,
        mark,
        "03",
        "Three questions",
        "Add a line on how the platform is organised.",
    )
    columns = (
        ("Discover", "Which startups deserve our attention?"),
        ("Evaluate", "Why are they promising, and what risks remain?"),
        ("Act", "Which funding, investors, or partners could help them grow?"),
    )
    gap = 0.2
    card_w = (CONTENT_W - gap * 2) / 3

    for index, (title, question) in enumerate(columns):
        x = ML + index * (card_w + gap)
        shape = card(
            slide,
            x,
            1.5,
            card_w,
            5.38,
        )
        shape.name = title
        pill(
            slide,
            x + 0.24,
            1.74,
            f"{index + 1:02d}",
        )
        textbox(
            slide,
            x + 0.7,
            1.8,
            card_w - 0.92,
            0.4,
            [
                {
                    "text": title,
                    "size": 20,
                    "bold": True,
                    "color": NAVY,
                    "font": DISPLAY,
                },
            ],
            ml=0,
            mr=0,
            mt=0,
            mb=0,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        textbox(
            slide,
            x + 0.16,
            2.4,
            card_w - 0.32,
            1.25,
            [
                {
                    "text": question,
                    "size": 15,
                    "color": INK,
                    "line": 1.15,
                },
            ],
            ml=0.1,
            mr=0.1,
            mt=0,
            mb=0,
        )
        divider = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(x + 0.28),
            Inches(3.72),
            Inches(0.7),
            Inches(0.025),
        )
        fill_rgb(divider, TEAL)
        no_line(divider)
        body = textbox(
            slide,
            x + 0.16,
            3.95,
            card_w - 0.32,
            2.6,
            [
                {
                    "text": "Write how the product answers this.",
                    "size": 15,
                    "italic": True,
                    "color": HINT,
                    "line": 1.25,
                },
            ],
            ml=0.1,
            mr=0.1,
            mt=0,
            mb=0,
        )
        body.name = f"{title} answer"

    footer(slide, "03  /  06")
    notes(
        slide,
        "Platform frame. The three questions can stay. Replace each answer underneath.",
    )


def product(prs: Presentation, mark: io.BytesIO) -> None:
    slide = blank_slide(prs)
    header(
        slide,
        mark,
        "04",
        "The product",
        "Add a line on what someone does in the product.",
    )
    stage = rounded(
        slide,
        ML,
        1.5,
        7.35,
        5.38,
        radius=0.06,
    )
    fill_rgb(stage, SOFT)
    stage.line.color.rgb = TEAL
    stage.line.width = Pt(1.25)
    line = stage.line._get_or_add_ln()
    etree.SubElement(
        line,
        qn("a:prstDash"),
        {"val": "dash"},
    )
    stage.name = "Product visual"
    textbox(
        slide,
        ML + 0.4,
        3.35,
        6.55,
        1.3,
        [
            {
                "text": "PRODUCT VISUAL",
                "size": 12,
                "bold": True,
                "color": TEAL,
                "font": DISPLAY,
                "align": PP_ALIGN.CENTER,
                "spc": 180,
            },
            {
                "text": "Drop a screenshot here",
                "size": 18,
                "italic": True,
                "color": HINT,
                "align": PP_ALIGN.CENTER,
                "before": 8,
            },
        ],
        ml=0,
        mr=0,
        mt=0,
        mb=0,
    )
    rows = (
        ("What it does", "One sentence on the job it performs."),
        ("Who uses it", "The person and the moment they open it."),
        ("What they see", "The decision the screen makes obvious."),
        ("What happens next", "The action it helps them take."),
    )
    row_x = ML + 7.35 + 0.2
    row_w = W - ML - row_x
    row_h = 1.22
    row_gap = 0.16

    for index, (label, prompt) in enumerate(rows):
        y = 1.5 + index * (row_h + row_gap)
        shape = card(
            slide,
            row_x,
            y,
            row_w,
            row_h,
            radius=0.12,
        )
        shape.name = label
        textbox(
            slide,
            row_x + 0.16,
            y + 0.14,
            row_w - 0.32,
            0.28,
            [
                {
                    "text": label.upper(),
                    "size": 11,
                    "bold": True,
                    "color": TEAL,
                    "font": DISPLAY,
                    "spc": 120,
                },
            ],
            ml=0.06,
            mr=0,
            mt=0,
            mb=0,
        )
        textbox(
            slide,
            row_x + 0.16,
            y + 0.46,
            row_w - 0.32,
            0.64,
            [
                {
                    "text": prompt,
                    "size": 14,
                    "italic": True,
                    "color": HINT,
                    "line": 1.15,
                },
            ],
            ml=0.06,
            mr=0.04,
            mt=0,
            mb=0,
        )

    footer(slide, "04  /  06")
    notes(
        slide,
        "Product. Paste a screenshot over the dashed frame. Replace the four notes on the right.",
    )


def vision(prs: Presentation, mark: io.BytesIO) -> None:
    slide = blank_slide(prs)
    header(
        slide,
        mark,
        "05",
        "The vision",
        "Add a line that sets up the close.",
    )
    quote = card(
        slide,
        ML,
        1.5,
        CONTENT_W,
        2.05,
        radius=0.08,
    )
    quote.name = "Vision statement"
    accent = rounded(
        slide,
        ML + 0.28,
        1.74,
        0.07,
        1.56,
        radius=0.5,
    )
    gradient(
        accent,
        TEAL_BRIGHT,
        TEAL_DEEP,
        angle=90,
    )
    textbox(
        slide,
        ML + 0.52,
        1.68,
        CONTENT_W - 0.8,
        1.7,
        [
            {
                "text": "CLOSING LINE",
                "size": 11,
                "bold": True,
                "color": TEAL,
                "font": DISPLAY,
                "spc": 160,
            },
            {
                "text": "Write the vision in one or two sentences.",
                "size": 22,
                "italic": True,
                "color": HINT,
                "before": 8,
                "line": 1.15,
            },
        ],
        ml=0.08,
        mr=0.1,
        mt=0,
        mb=0,
    )
    titles = ("Who it serves", "What changes", "The ask")
    prompts = (
        "Funds, accelerators, or partners — name them.",
        "What gets faster, clearer, or funded.",
        "What you want from the room.",
    )
    gap = 0.2
    card_w = (CONTENT_W - gap * 2) / 3

    for index, (title, prompt) in enumerate(zip(titles, prompts)):
        x = ML + index * (card_w + gap)
        shape = card(
            slide,
            x,
            3.75,
            card_w,
            3.13,
        )
        shape.name = title
        textbox(
            slide,
            x + 0.16,
            3.96,
            card_w - 0.32,
            0.4,
            [
                {
                    "text": title,
                    "size": 16,
                    "bold": True,
                    "color": NAVY,
                    "font": DISPLAY,
                },
            ],
            ml=0.08,
            mr=0.06,
            mt=0,
            mb=0,
        )
        textbox(
            slide,
            x + 0.16,
            4.46,
            card_w - 0.32,
            2.15,
            [
                {
                    "text": prompt,
                    "size": 14,
                    "italic": True,
                    "color": HINT,
                    "line": 1.2,
                },
            ],
            ml=0.08,
            mr=0.06,
            mt=0,
            mb=0,
        )

    footer(slide, "05  /  06")
    notes(
        slide,
        "Close. Replace the vision line, then who it serves, what changes, and the ask.",
    )


def closing_background() -> io.BytesIO:
    """Fit the reef still to the slide and shade the lower half for type."""
    image = Image.open(CLOSING).convert("RGB")
    width, height = image.size
    target = W / H

    if width / height > target:
        crop_w = int(height * target)
        left = (width - crop_w) // 2
        image = image.crop((left, 0, left + crop_w, height))
    else:
        crop_h = int(width / target)
        top = max(0, int((height - crop_h) * 0.2))
        image = image.crop((0, top, width, top + crop_h))

    image = image.resize((1920, 1080), Image.Resampling.LANCZOS)
    image = ImageEnhance.Brightness(image).enhance(0.9)

    alpha = Image.new("L", (1, 1080))

    for y in range(1080):
        fade = (y / 1080 - 0.36) / 0.64
        if fade < 0:
            value = 0
        else:
            value = int(min(1.0, fade) ** 1.08 * 175)
        alpha.putpixel((0, y), value)

    alpha = alpha.resize((1920, 1080))
    scrim = Image.new(
        "RGBA",
        (1920, 1080),
        (6, 32, 48, 255),
    )
    scrim.putalpha(alpha)
    framed = image.convert("RGBA")
    framed.alpha_composite(scrim)
    buffer = io.BytesIO()
    framed.convert("RGB").save(
        buffer,
        format="JPEG",
        quality=90,
    )
    buffer.seek(0)

    return buffer


def thanks(prs: Presentation, lockup: io.BytesIO) -> None:
    """Close on the reef, the logo, and a thank-you to the room."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(
        closing_background(),
        0,
        0,
        Inches(W),
        Inches(H),
    )
    lockup.seek(0)
    logo_w = 5.2
    logo_h = logo_w / lockup.aspect
    pad_x = 0.34
    pad_y = 0.24
    card_w = logo_w + pad_x * 2
    card_h = logo_h + pad_y * 2
    card_x = (W - card_w) / 2
    card_y = 0.78
    plate = rounded(
        slide,
        card_x,
        card_y,
        card_w,
        card_h,
        radius=0.1,
    )
    fill_rgb(plate, WHITE)
    no_line(plate)
    add_shadow(
        plate,
        alpha=32000,
        blur=560000,
        dist=100000,
    )
    slide.shapes.add_picture(
        lockup,
        Inches(card_x + pad_x),
        Inches(card_y + pad_y),
        Inches(logo_w),
        Inches(logo_h),
    )
    textbox(
        slide,
        1.4,
        card_y + card_h + 0.38,
        W - 2.8,
        1.45,
        [
            {
                "text": "Thank you",
                "size": 46,
                "bold": True,
                "color": WHITE,
                "font": DISPLAY,
                "align": PP_ALIGN.CENTER,
            },
            {
                "text": "for watching",
                "size": 22,
                "italic": True,
                "color": MIST,
                "align": PP_ALIGN.CENTER,
                "before": 4,
            },
        ],
        ml=0,
        mr=0,
        mt=0,
        mb=0,
    )
    textbox(
        slide,
        W - ML - 1.3,
        7.12,
        1.3,
        0.24,
        [
            {
                "text": "06  /  06",
                "size": 11,
                "bold": True,
                "color": MIST,
                "font": DISPLAY,
                "align": PP_ALIGN.RIGHT,
                "spc": 80,
            },
        ],
        ml=0,
        mr=0,
        mt=0,
        mb=0,
    )
    notes(
        slide,
        "Ending. Thank the people who stayed, then stop.",
    )


def build() -> Path:
    """Write the six-slide template and return its path."""
    lockup = cropped_logo()
    mark = icon_mark()
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    prs.core_properties.title = "Ocean Opportunity Radar"
    prs.core_properties.subject = "Pitch deck template"
    prs.core_properties.category = "A’Pelago × Redstone"
    cover(prs, lockup)
    problem(prs, mark)
    questions(prs, mark)
    product(prs, mark)
    vision(prs, mark)
    thanks(prs, lockup)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)

    return OUT


if __name__ == "__main__":
    print(build())
